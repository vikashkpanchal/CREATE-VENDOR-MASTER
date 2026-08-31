"""Excel export of vendor records.

Every vendor's emails stay exactly as entered - one semicolon-separated
string in a single "Vendor Email ID" column. They are deliberately NOT
split across "Vendor Email ID 1/2/3..." columns, so an exported sheet can
be edited and re-imported without the address list changing shape.
"""

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from vendor_app.config import KEYS, LABELS

HEADER_FILL = "1F6AA5"
HEADER_FONT_COLOR = "FFFFFF"
SHEET_NAME = "Vendor Master"


def build_export_dataframe(records: list) -> pd.DataFrame:
    """One row per vendor, one column per field, in the mandated order.

    "Vendor Email ID" holds the vendor's full semicolon-separated address
    list verbatim - no fanning out into numbered columns.
    """
    columns = [
        "Sr. No.",
        LABELS["vendor_code"],
        LABELS["vendor_name"],
        LABELS["vendor_email"],
        LABELS["vendor_owner_name"],
        LABELS["vendor_owner_contact"],
        LABELS["vendor_owner_email"],
        LABELS["vendor_supervisor_name"],
        LABELS["vendor_supervisor_contact"],
        LABELS["vendor_supervisor_email"],
        LABELS["vendor_type"],
        LABELS["city"],
        LABELS["state"],
        "Status",
    ]

    rows = []
    for i, record in enumerate(records, start=1):
        row = {"Sr. No.": i}
        for key in KEYS:
            row[LABELS[key]] = record.get(key, "")
        row["Status"] = record.get("status", "Active")
        rows.append(row)

    return pd.DataFrame(rows, columns=columns)


def _style_worksheet(worksheet, columns):
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_font = Font(bold=True, color=HEADER_FONT_COLOR)
    wrap_center = Alignment(wrap_text=True, vertical="center", horizontal="center")

    for col_idx, col_name in enumerate(columns, start=1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = wrap_center

        data_lengths = [
            len(str(worksheet.cell(row=r, column=col_idx).value or ""))
            for r in range(2, worksheet.max_row + 1)
        ]
        max_len = max([len(col_name)] + data_lengths)
        worksheet.column_dimensions[cell.column_letter].width = min(max(max_len + 4, 14), 42)

    worksheet.freeze_panes = "A2"
    worksheet.row_dimensions[1].height = 32


def export_records_to_excel(records: list, path: str) -> str:
    """Write `records` to a formatted .xlsx workbook at `path`. Returns path."""
    df = build_export_dataframe(records)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=SHEET_NAME)
        _style_worksheet(writer.sheets[SHEET_NAME], list(df.columns))
    return path


def export_audit_log_to_excel(entries: list, path: str, columns=None, headers=None) -> str:
    """Write change-log `entries` (most-recent-first dicts, see audit.py) to a
    formatted .xlsx workbook at `path`. `columns`/`headers` let the equipment
    log export with its own column set. Returns path."""
    from vendor_app.config import AUDIT_COLUMNS, AUDIT_WRAPPED_LABELS

    columns = list(columns or AUDIT_COLUMNS)
    header_map = headers or AUDIT_WRAPPED_LABELS
    headers = [header_map.get(col, col).replace("\n", " ") for col in columns]
    rows = [{headers[i]: entry.get(col, "") for i, col in enumerate(columns)} for entry in entries]
    df = pd.DataFrame(rows, columns=headers)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Audit Log")
        _style_worksheet(writer.sheets["Audit Log"], headers)
    return path


def export_equipment_to_excel(records: list, path: str) -> str:
    """Write equipment master `records` to a formatted .xlsx at `path`."""
    from vendor_app.config import EQUIPMENT_KEYS, EQUIPMENT_LABELS

    columns = ["Sr. No."] + [EQUIPMENT_LABELS[k] for k in EQUIPMENT_KEYS]
    rows = []
    for i, record in enumerate(records, start=1):
        row = {"Sr. No.": i}
        for key in EQUIPMENT_KEYS:
            row[EQUIPMENT_LABELS[key]] = record.get(key, "")
        rows.append(row)

    df = pd.DataFrame(rows, columns=columns)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Equipment Master")
        _style_worksheet(writer.sheets["Equipment Master"], columns)
    return path


def _export_simple(records: list, path: str, columns_keys, labels, sheet_name) -> str:
    """Shared writer for the ARC/FO/line-item sheets: Sr. No. then one column
    per key, in the order the grid shows them."""
    columns = ["Sr. No."] + [labels[k] for k in columns_keys]
    rows = []
    for i, record in enumerate(records, start=1):
        row = {"Sr. No.": i}
        for key in columns_keys:
            row[labels[key]] = record.get(key, "")
        rows.append(row)

    df = pd.DataFrame(rows, columns=columns)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        _style_worksheet(writer.sheets[sheet_name], columns)
    return path


def export_arcs_to_excel(records: list, path: str) -> str:
    """Table 1 as shown, including the per-contract derived columns."""
    from vendor_app.config import ARC_DERIVED_LABELS, ARC_DISPLAY_COLUMNS, ARC_LABELS

    labels = dict(ARC_LABELS)
    labels.update(ARC_DERIVED_LABELS)
    return _export_simple(records, path, ARC_DISPLAY_COLUMNS, labels, "ARC Data")


def export_fos_to_excel(records: list, path: str) -> str:
    """Table 2 as shown, including each frame order's item totals."""
    from vendor_app.config import FO_DERIVED_LABELS, FO_DISPLAY_COLUMNS, FO_LABELS

    labels = dict(FO_LABELS)
    labels.update(FO_DERIVED_LABELS)
    return _export_simple(records, path, FO_DISPLAY_COLUMNS, labels, "FO Data")


def export_report_to_excel(report, path: str) -> str:
    """Write one ARC/FO analysis report (see arc_analytics.Report) to .xlsx."""
    headers = [report.headers(wrapped=False)[key] for key in report.columns]
    rows = [dict(zip(headers, report.values(row))) for row in report.rows]
    df = pd.DataFrame(rows, columns=headers)

    sheet = report.title[:31]
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet)
        _style_worksheet(writer.sheets[sheet], headers)
    return path


def export_arc_analysis_to_excel(analysis, path: str) -> str:
    """The whole ARC & FO dashboard as one workbook.

    Sheet 1 is the 19 headline analyses; every table on the dashboard then
    follows on its own sheet, in the order the dashboard shows them, so the
    workbook reads the same way the screen does.
    """
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary = pd.DataFrame(analysis.summary_rows(), columns=["Analysis", "Value"])
        summary.to_excel(writer, index=False, sheet_name="Summary")
        _style_worksheet(writer.sheets["Summary"], ["Analysis", "Value"])

        used = {"Summary"}
        for report in analysis.full_reports():
            headers = [report.headers(wrapped=False)[key] for key in report.columns]
            rows = [dict(zip(headers, report.values(row))) for row in report.rows]
            # An empty report still gets its sheet: "nothing is expiring" is
            # an answer, and a missing tab reads as a missing analysis.
            frame = pd.DataFrame(rows, columns=headers)

            sheet = report.title[:31]
            suffix = 2
            while sheet in used:
                sheet = f"{report.title[:28]} {suffix}"
                suffix += 1
            used.add(sheet)

            frame.to_excel(writer, index=False, sheet_name=sheet)
            _style_worksheet(writer.sheets[sheet], headers)
    return path


def export_arc_value_to_excel(lines: list, path: str, title: str = "") -> str:
    """The priced annexure, as two sheets of live formulas.

    **Annexure 2** is the calculation, one or two rows per equipment
    category. **Annexure 1** summarises it, one row per contract, and reads
    its Impact straight out of Annexure 2's subtotal cell.

    Every multiplication and every total is written as an Excel FORMULA, not
    as a number this app worked out: the sheet can be audited cell by cell,
    a rate can be corrected in place and everything above and below it
    follows. The totals use SUBTOTAL(9,...), which ignores the nested
    subtotals inside its own range and re-totals whatever a filter leaves
    visible - which is how the source workbook does it.

    Calibri 10 throughout, headers and totals bold, total rows on a very
    light blue.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from vendor_app.arc_value import LINE_ARC_TOTAL, LINE_GRAND_TOTAL, LINE_OT
    from vendor_app.config import (
        ARC_SUMMARY_COLUMNS, ARC_SUMMARY_LABELS, ARC_VALUE_COLUMNS, ARC_VALUE_LABELS,
        OT_HOURS_PER_DAY, WORKING_DAYS_PER_MONTH,
    )

    BODY_FONT = Font(name="Calibri", size=10)
    BOLD_FONT = Font(name="Calibri", size=10, bold=True)
    # Black on white, bordered and bold: this annexure is printed and signed,
    # and a banded header wastes toner without adding anything the bold and
    # the border do not already say.
    HEAD_FONT = Font(name="Calibri", size=10, bold=True, color="000000")
    TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")   # very light blue
    HEAD_FILL = PatternFill("solid", fgColor="FFFFFF")
    thin = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
    MONEY = "#,##0"
    RATE = "#,##0.00"

    detail_numeric = {"eqp_qty", "qty", "monthly_rate", "value"}
    summary_numeric = {"existing_value", "revised_value", "impact"}

    def write_header(sheet, columns, labels, row):
        for index, key in enumerate(columns, start=1):
            cell = sheet.cell(row=row, column=index, value=labels[key])
            cell.fill = HEAD_FILL
            cell.font = HEAD_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="center",
                                       horizontal="center")
            cell.border = BORDER
        sheet.row_dimensions[row].height = 32

    def fit(sheet, columns, labels, header_row, last_row):
        for index, key in enumerate(columns, start=1):
            letter = get_column_letter(index)
            longest = max(
                [len(labels[key])]
                + [len(str(sheet.cell(row=r, column=index).value or ""))
                   for r in range(header_row + 1, last_row + 1)]
            )
            sheet.column_dimensions[letter].width = min(max(longest + 3, 11), 42)
        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    book = Workbook()

    # ------------------------------------------------------- Annexure 2 --
    detail = book.active
    detail.title = "Annexure 2"
    row_index = 1
    if title:
        cell = detail.cell(row=1, column=1, value=title)
        cell.font = Font(name="Calibri", size=12, bold=True)
        row_index = 3
    head2 = row_index
    write_header(detail, ARC_VALUE_COLUMNS, ARC_VALUE_LABELS, head2)

    column_of = {key: i + 1 for i, key in enumerate(ARC_VALUE_COLUMNS)}
    L = {key: get_column_letter(index) for key, index in column_of.items()}
    first_data = head2 + 1
    last_data = head2 + len(lines)

    block_start = first_data          # first row of the contract being written
    mcm_row = None                    # the MCM row an OT row multiplies out of
    summaries = []                    # (excel row of the subtotal, summary dict)

    for offset, line in enumerate(lines):
        row = first_data + offset
        kind = line["kind"]
        is_total = kind in (LINE_ARC_TOTAL, LINE_GRAND_TOTAL)

        for key in ARC_VALUE_COLUMNS:
            value = line.get(key, "")
            if key in detail_numeric and value not in ("", None):
                value = float(value)
            cell = detail.cell(row=row, column=column_of[key], value=value)
            cell.border = BORDER
            cell.font = BOLD_FONT if is_total else BODY_FONT
            if is_total:
                cell.fill = TOTAL_FILL
            if key in detail_numeric:
                cell.number_format = RATE if key == "monthly_rate" else MONEY
                cell.alignment = Alignment(horizontal="right")

        if kind == LINE_ARC_TOTAL:
            detail[f"{L['eqp_qty']}{row}"] = (
                f"=SUBTOTAL(9,{L['eqp_qty']}{block_start}:{L['eqp_qty']}{row - 1})"
            )
            detail[f"{L['value']}{row}"] = (
                f"=SUBTOTAL(9,{L['value']}{block_start}:{L['value']}{row - 1})"
            )
            if line.get("summary"):
                summaries.append((row, line["summary"]))
            block_start = row + 1
            mcm_row = None
        elif kind == LINE_GRAND_TOTAL:
            # SUBTOTAL ignores the nested SUBTOTALs in its own range, so this
            # sums the priced rows once rather than double-counting them.
            detail[f"{L['eqp_qty']}{row}"] = (
                f"=SUBTOTAL(9,{L['eqp_qty']}{first_data}:{L['eqp_qty']}{row - 1})"
            )
            detail[f"{L['value']}{row}"] = (
                f"=SUBTOTAL(9,{L['value']}{first_data}:{L['value']}{row - 1})"
            )
        elif kind == LINE_OT:
            hours = OT_HOURS_PER_DAY.get(str(line.get("working_shift", "")), 0)
            if mcm_row is not None:
                detail[f"{L['qty']}{row}"] = (
                    f"={L['eqp_qty']}{mcm_row}*{L['qty']}{mcm_row}"
                    f"*{WORKING_DAYS_PER_MONTH}*{hours}"
                )
            # An OT line's quantity already carries the equipment count, so
            # its value is quantity x rate and nothing else.
            detail[f"{L['value']}{row}"] = (
                f"={L['qty']}{row}*{L['monthly_rate']}{row}"
            )
        else:
            mcm_row = row
            detail[f"{L['value']}{row}"] = (
                f"={L['eqp_qty']}{row}*{L['qty']}{row}*{L['monthly_rate']}{row}"
            )
            detail.cell(row=row, column=column_of["value"]).number_format = MONEY

    for key in ("eqp_qty", "qty", "monthly_rate", "value"):
        for row in range(first_data, last_data + 1):
            cell = detail.cell(row=row, column=column_of[key])
            cell.number_format = RATE if key == "monthly_rate" else MONEY
            cell.alignment = Alignment(horizontal="right")

    fit(detail, ARC_VALUE_COLUMNS, ARC_VALUE_LABELS, head2, last_data)
    if lines:
        detail.auto_filter.ref = (
            f"A{head2}:{get_column_letter(len(ARC_VALUE_COLUMNS))}{last_data}"
        )

    # ------------------------------------------------------- Annexure 1 --
    summary = book.create_sheet("Annexure 1", 0)
    cell = summary.cell(row=1, column=1,
                        value=title or "Contract Amendment Value Calculation")
    cell.font = Font(name="Calibri", size=12, bold=True)
    summary.cell(row=2, column=1, value="Summary of Annexure 2 - one row per ARC"
                 ).font = Font(name="Calibri", size=10, italic=True)
    head1 = 4
    write_header(summary, ARC_SUMMARY_COLUMNS, ARC_SUMMARY_LABELS, head1)

    scol = {key: i + 1 for i, key in enumerate(ARC_SUMMARY_COLUMNS)}
    S = {key: get_column_letter(index) for key, index in scol.items()}
    for offset, (detail_row, entry) in enumerate(summaries, start=1):
        row = head1 + offset
        for key in ARC_SUMMARY_COLUMNS:
            value = entry.get(key, "")
            if key in summary_numeric and value not in ("", None):
                value = float(value)
            cell = summary.cell(row=row, column=scol[key], value=value)
            cell.border = BORDER
            cell.font = BODY_FONT
            if key in summary_numeric:
                cell.number_format = MONEY
                cell.alignment = Alignment(horizontal="right")
        # Impact is READ from Annexure 2's own subtotal, and the revised
        # value is derived from it here, so the two sheets cannot drift.
        summary[f"{S['impact']}{row}"] = f"='Annexure 2'!{L['value']}{detail_row}"
        summary[f"{S['revised_value']}{row}"] = (
            f"={S['existing_value']}{row}+{S['impact']}{row}"
        )

    if summaries:
        total_row = head1 + len(summaries) + 1
        summary.cell(row=total_row, column=scol["vendor_name"], value="Total")
        for key in summary_numeric:
            letter = S[key]
            summary[f"{letter}{total_row}"] = (
                f"=SUBTOTAL(9,{letter}{head1 + 1}:{letter}{total_row - 1})"
            )
        for key in ARC_SUMMARY_COLUMNS:
            cell = summary.cell(row=total_row, column=scol[key])
            cell.font = BOLD_FONT
            cell.fill = TOTAL_FILL
            cell.border = BORDER
            if key in summary_numeric:
                cell.number_format = MONEY
                cell.alignment = Alignment(horizontal="right")
        fit(summary, ARC_SUMMARY_COLUMNS, ARC_SUMMARY_LABELS, head1, total_row)

    book.save(path)
    return path


# --------------------------------------------- all four masters, one file --
# One sheet per master, each with the columns that master already exports, in
# the same order. The names are what the importer looks for when the file
# comes back, so they are constants rather than literals typed twice.
MASTER_SHEETS = {
    "vendors": "Vendor Master",
    "equipment": "Equipment Master",
    "arcs": "ARC Master",
    "fos": "FO Master",
}


def _master_frames(vendors, equipment, arcs, fos):
    """(sheet name, DataFrame) for each master, in its own column order."""
    from vendor_app.config import (
        ARC_DERIVED_LABELS, ARC_DISPLAY_COLUMNS, ARC_LABELS, EQUIPMENT_KEYS,
        EQUIPMENT_LABELS, FO_DERIVED_LABELS, FO_DISPLAY_COLUMNS, FO_LABELS,
    )

    def simple(records, keys, labels):
        columns = ["Sr. No."] + [labels[k] for k in keys]
        rows = []
        for index, record in enumerate(records, start=1):
            row = {"Sr. No.": index}
            for key in keys:
                row[labels[key]] = record.get(key, "")
            rows.append(row)
        return pd.DataFrame(rows, columns=columns)

    arc_labels = dict(ARC_LABELS); arc_labels.update(ARC_DERIVED_LABELS)
    fo_labels = dict(FO_LABELS); fo_labels.update(FO_DERIVED_LABELS)
    return [
        (MASTER_SHEETS["vendors"], build_export_dataframe(vendors)),
        (MASTER_SHEETS["equipment"],
         simple(equipment, EQUIPMENT_KEYS, EQUIPMENT_LABELS)),
        (MASTER_SHEETS["arcs"], simple(arcs, ARC_DISPLAY_COLUMNS, arc_labels)),
        (MASTER_SHEETS["fos"], simple(fos, FO_DISPLAY_COLUMNS, fo_labels)),
    ]


def export_all_masters_to_excel(vendors, equipment, arcs, fos, path: str) -> dict:
    """Every master in one workbook - one sheet each, arrangement unchanged.

    Each sheet is exactly what that master's own export produces, so this is
    a backup that can be read by eye, edited, and handed straight back to
    the one-click import. Returns the row count per sheet.
    """
    frames = _master_frames(vendors, equipment, arcs, fos)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet, frame in frames:
            frame.to_excel(writer, index=False, sheet_name=sheet)
            _style_worksheet(writer.sheets[sheet], list(frame.columns))
    return {sheet: len(frame) for sheet, frame in frames}
